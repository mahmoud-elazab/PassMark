"""
Writes processing results to an Excel file.
Layout matches the FT Data Summary Report format.
Dynamic columns based on selected operating modes.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import column_index_from_string


def apply_outer_border(ws, cell_range, style="medium"):
    """Apply a border only around the outside of a cell range."""
    start_cell, end_cell = cell_range.split(":")
    start_col_letters = "".join([c for c in start_cell if c.isalpha()])
    start_row = int("".join([c for c in start_cell if c.isdigit()]))
    end_col_letters = "".join([c for c in end_cell if c.isalpha()])
    end_row = int("".join([c for c in end_cell if c.isdigit()]))

    start_col = column_index_from_string(start_col_letters)
    end_col = column_index_from_string(end_col_letters)

    side = Side(style=style)

    for r in range(start_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            cell = ws.cell(row=r, column=c)
            b = cell.border
            left = side if c == start_col else b.left
            right = side if c == end_col else b.right
            top = side if r == start_row else b.top
            bottom = side if r == end_row else b.bottom
            cell.border = Border(left=left, right=right, top=top, bottom=bottom)


def is_out_of_tolerance(mode, port, avg_value, mode_definitions):
    """Check if a voltage value is outside tolerance for a given mode/port."""
    if avg_value in ("", None):
        return False
    try:
        v = float(avg_value)
    except (TypeError, ValueError):
        return False

    mode_def = mode_definitions.get(mode)
    if mode_def is None:
        return False

    port_def = mode_def.get(f'port{port}')
    if port_def is None:
        return False

    nominal = port_def['voltage']
    tol = 0.25

    return not (nominal - tol <= v <= nominal + tol)


def get_column_letter(col_index):
    """Convert column number to Excel letter (1=A, 2=B, etc.)."""
    result = ""
    while col_index > 0:
        col_index, remainder = divmod(col_index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def write_results_to_excel(results_data, output_path, stage_label="Initial", reference_currents=None):
    """
    Write results to an Excel file matching the FT Data Summary Report layout.
    Layout is dynamic based on selected operating modes.

    Args:
        results_data: Dictionary returned by process_files()
        output_path: Where to save the .xlsx file
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "FT Data Summary"

    # Styles
    title_font = Font(bold=True, size=14)
    header_font = Font(bold=True, size=11)
    data_font = Font(size=11)
    criteria_font = Font(bold=True, size=11)

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    bold_border = Border(
        left=Side(style="medium"),
        right=Side(style="medium"),
        top=Side(style="medium"),
        bottom=Side(style="medium")
    )

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    green_fill = PatternFill(
        start_color="B5E6A2", end_color="B5E6A2", fill_type="solid"
    )
    green_font = Font(bold=True, size=11)

    black_fill = PatternFill(
        start_color="000000", end_color="000000", fill_type="solid"
    )

    grey_fill = PatternFill(
        start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"
    )

    bad_number_fill = PatternFill(
        start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"
    )
    bad_number_font = Font(color="9C0006")

    # Data from results
    results = results_data['results']
    polarity = results_data.get('reverse_polarity', {})
    selected_modes = sorted(results_data.get('selected_modes', [5, 6, 7]))
    mode_definitions = results_data.get('mode_definitions', {})

    # Group results by device
    devices = {}
    device_order = []
    for r in results:
        device = r['device']
        if device not in devices:
            devices[device] = {}
            device_order.append(device)
        key = (r['mode'], r['port'])
        devices[device][key] = r['average']

    # Optional: warnings if reference Excel has fewer rows than output
    if reference_currents:
        out_n = len(device_order)
        ref_n = reference_currents.get("source_rows", 0)
        if ref_n < out_n:
            results_data.setdefault("warnings", []).append(
                f"Reference Excel has {ref_n} data row(s) but output needs {out_n}. "
                f"Remaining current cells left blank."
            )

    # Build column map dynamically
    col_map = {}
    loaded_current_cols = {}

    current_col = 3
    for mode in selected_modes:
        col_map[(mode, 1)] = current_col
        col_map[(mode, 2)] = current_col + 1
        loaded_current_cols[mode] = current_col + 2
        current_col += 3

    led_col = current_col
    rev_pol_col = current_col + 1
    unloaded_col = current_col + 2
    last_col = current_col + 2

    # Column widths
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 66 / 7

    for mode in selected_modes:
        ws.column_dimensions[get_column_letter(col_map[(mode, 1)])].width = 72 / 7
        ws.column_dimensions[get_column_letter(col_map[(mode, 2)])].width = 72 / 7
        ws.column_dimensions[get_column_letter(loaded_current_cols[mode])].width = 72 / 7

    ws.column_dimensions[get_column_letter(led_col)].width = 72 / 7
    ws.column_dimensions[get_column_letter(rev_pol_col)].width = 72 / 7
    ws.column_dimensions[get_column_letter(unloaded_col)].width = 72 / 7

    # Row heights
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 16
    ws.row_dimensions[3].height = 16
    ws.row_dimensions[4].height = 16
    ws.row_dimensions[5].height = 49

    # Shortcut
    def cl(idx):
        return get_column_letter(idx)

    # Row 2: Title (centered, green)
    ws.merge_cells(f'B2:{cl(last_col)}2')
    cell = ws['B2']
    cell.value = f"{stage_label} - @ 25° C"
    cell.font = title_font
    cell.fill = green_fill
    cell.alignment = center
    cell.border = bold_border
    for c in range(3, last_col + 1):
        ws.cell(row=2, column=c).border = bold_border

    # Row 3: Test: merged with row 4
    ws.merge_cells('B3:B4')
    cell = ws['B3']
    cell.value = "Test:"
    cell.font = green_font
    cell.fill = green_fill
    cell.alignment = center
    cell.border = bold_border
    ws['B4'].border = bold_border

    # Row 3: 13.5 V Input merged across all columns
    ws.merge_cells(f'C3:{cl(last_col)}3')
    cell = ws.cell(row=3, column=3)
    cell.value = "13.5 V Input"
    cell.font = green_font
    cell.fill = green_fill
    cell.alignment = center
    cell.border = bold_border
    for c in range(4, last_col + 1):
        ws.cell(row=3, column=c).border = bold_border

    # Row 4: Mode section headers
    for mode in selected_modes:
        mode_def = mode_definitions.get(mode, {})
        label = mode_def.get('label', f'Mode {mode}')
        start = col_map[(mode, 1)]
        end = loaded_current_cols[mode]

        ws.merge_cells(
            start_row=4, start_column=start,
            end_row=4, end_column=end
        )
        cell = ws.cell(row=4, column=start)
        cell.value = label
        cell.font = green_font
        cell.fill = green_fill
        cell.alignment = center
        cell.border = bold_border
        for c in range(start + 1, end + 1):
            ws.cell(row=4, column=c).border = bold_border

    # Row 4-5: LED On? merged
    ws.merge_cells(
        start_row=4, start_column=led_col,
        end_row=5, end_column=led_col
    )
    cell = ws.cell(row=4, column=led_col)
    cell.value = "LED On?"
    cell.font = green_font
    cell.fill = green_fill
    cell.alignment = center
    cell.border = bold_border
    ws.cell(row=5, column=led_col).border = bold_border

    # Row 4-5: Reverse Polarity merged
    ws.merge_cells(
        start_row=4, start_column=rev_pol_col,
        end_row=5, end_column=rev_pol_col
    )
    cell = ws.cell(row=4, column=rev_pol_col)
    cell.value = "Reverse polarity"
    cell.font = green_font
    cell.fill = green_fill
    cell.alignment = center
    cell.border = bold_border
    ws.cell(row=5, column=rev_pol_col).border = bold_border

    # Row 4-5: Unloaded Current merged
    ws.merge_cells(
        start_row=4, start_column=unloaded_col,
        end_row=5, end_column=unloaded_col
    )
    cell = ws.cell(row=4, column=unloaded_col)
    cell.value = "Unloaded Current (mA)"
    cell.font = green_font
    cell.fill = green_fill
    cell.alignment = center
    cell.border = bold_border
    ws.cell(row=5, column=unloaded_col).border = bold_border

    # Row 5: Sample # sub-header
    cell = ws.cell(row=5, column=2)
    cell.value = "Sample #"
    cell.font = header_font
    cell.fill = green_fill
    cell.alignment = center
    cell.border = bold_border

    # Row 5: Sub-headers for each mode
    for mode in selected_modes:
        mode_def = mode_definitions.get(mode, {})
        p1_col = col_map[(mode, 1)]
        p2_col = col_map[(mode, 2)]
        lc_col = loaded_current_cols[mode]

        port1_def = mode_def.get('port1')
        port2_def = mode_def.get('port2')

        p1_label = (
            f"Port 1 (V) {port1_def['watts']}W"
            if port1_def else "Port 1 (V) 0W"
        )
        p2_label = (
            f"Port 2 (V) {port2_def['watts']}W"
            if port2_def else "Port 2 (V) 0W"
        )

        cell = ws.cell(row=5, column=p1_col)
        cell.value = p1_label
        cell.font = header_font
        cell.fill = green_fill
        cell.alignment = center
        cell.border = bold_border

        cell = ws.cell(row=5, column=p2_col)
        cell.value = p2_label
        cell.font = header_font
        cell.fill = green_fill
        cell.alignment = center
        cell.border = bold_border

        cell = ws.cell(row=5, column=lc_col)
        cell.value = "Loaded Current (A)"
        cell.font = header_font
        cell.fill = green_fill
        cell.alignment = center
        cell.border = bold_border

    # Data rows
    row = 6
    for device in device_order:
        values = devices[device]
        ws.row_dimensions[row].height = 16

        # Sample # (as text to avoid number formatting)
        cell = ws.cell(row=row, column=2, value=str(device))
        cell.font = data_font
        cell.alignment = center
        cell.border = thin_border

        for mode in selected_modes:
            mode_def = mode_definitions.get(mode, {})
            p1_col = col_map[(mode, 1)]
            p2_col = col_map[(mode, 2)]
            lc_col = loaded_current_cols[mode]

            port1_def = mode_def.get('port1')
            port2_def = mode_def.get('port2')

            # Port 1
            if port1_def is not None:
                val = values.get((mode, 1), '')
                cell = ws.cell(row=row, column=p1_col)
                cell.value = val if val == '' else float(val)
                cell.number_format = '0.00'
                cell.font = data_font
                cell.alignment = center
                cell.border = thin_border
                if is_out_of_tolerance(mode, 1, val, mode_definitions):
                    cell.fill = bad_number_fill
                    cell.font = bad_number_font
            else:
                cell = ws.cell(row=row, column=p1_col)
                cell.fill = black_fill
                cell.border = thin_border

            # Port 2
            if port2_def is not None:
                val = values.get((mode, 2), '')
                cell = ws.cell(row=row, column=p2_col)
                cell.value = val if val == '' else float(val)
                cell.number_format = '0.00'
                cell.font = data_font
                cell.alignment = center
                cell.border = thin_border
                if is_out_of_tolerance(mode, 2, val, mode_definitions):
                    cell.fill = bad_number_fill
                    cell.font = bad_number_font
            else:
                cell = ws.cell(row=row, column=p2_col)
                cell.fill = black_fill
                cell.border = thin_border

            # Loaded Current (optional from reference Excel)
            loaded_val = None
            if reference_currents:
                loaded_by_mode = reference_currents.get("loaded_A_by_mode", {})
                # row_index is 0-based position within device_order
                row_index = row - 6
                mode_list = loaded_by_mode.get(mode, [])
                if 0 <= row_index < len(mode_list):
                    loaded_val = mode_list[row_index]

            cell = ws.cell(row=row, column=lc_col)
            cell.value = "" if loaded_val is None else float(loaded_val)
            cell.number_format = "0.00" if loaded_val is not None else "General"
            cell.font = data_font
            cell.alignment = center
            cell.border = thin_border

        # LED On?
        cell = ws.cell(row=row, column=led_col)
        cell.value = 'On'
        cell.font = data_font
        cell.alignment = center
        cell.border = thin_border


        # Reverse Polarity
        cell = ws.cell(row=row, column=rev_pol_col)
        cell.value = polarity.get(device, '')
        cell.font = data_font
        cell.alignment = center
        cell.border = thin_border

        # Unloaded Current (optional from reference Excel)
        unloaded_val = None
        if reference_currents:
            row_index = row - 6  # 0-based data-row index
            ul = reference_currents.get("unloaded_mA", [])
            if 0 <= row_index < len(ul):
                unloaded_val = ul[row_index]

        cell = ws.cell(row=row, column=unloaded_col)
        cell.value = "" if unloaded_val is None else float(unloaded_val)
        cell.number_format = "0.00" if unloaded_val is not None else "General"
        cell.font = data_font
        cell.alignment = center
        cell.border = thin_border

        row += 1

    # Criteria row
    criteria_row = row
    ws.row_dimensions[criteria_row].height = 24

    # Sample # criteria label
    cell = ws.cell(row=criteria_row, column=2, value="Criteria")
    cell.font = criteria_font
    cell.alignment = center
    cell.border = bold_border

    # Criteria values for each mode
    for mode in selected_modes:
        mode_def = mode_definitions.get(mode, {})
        p1_col = col_map[(mode, 1)]
        p2_col = col_map[(mode, 2)]
        lc_col = loaded_current_cols[mode]

        port1_def = mode_def.get('port1')
        port2_def = mode_def.get('port2')

        # Port 1 criteria
        if port1_def is not None:
            ws.merge_cells(
                start_row=criteria_row, start_column=p1_col,
                end_row=criteria_row, end_column=p1_col
            )
            cell = ws.cell(row=criteria_row, column=p1_col)
            cell.value = f"{port1_def['voltage']}±0.25V"
            cell.font = criteria_font
            cell.alignment = center
            cell.border = bold_border
        else:
            cell = ws.cell(row=criteria_row, column=p1_col)
            cell.fill = black_fill
            cell.border = bold_border

        # Port 2 criteria
        if port2_def is not None:
            cell = ws.cell(row=criteria_row, column=p2_col)
            cell.value = f"{port2_def['voltage']}±0.25V"
            cell.font = criteria_font
            cell.alignment = center
            cell.border = bold_border
        else:
            cell = ws.cell(row=criteria_row, column=p2_col)
            cell.fill = black_fill
            cell.border = bold_border

        # Loaded Current criteria (grey)
        cell = ws.cell(row=criteria_row, column=lc_col)
        cell.fill = grey_fill
        cell.border = bold_border

    # LED On? criteria
    cell = ws.cell(row=criteria_row, column=led_col)
    cell.value = "On"
    cell.font = criteria_font
    cell.alignment = center
    cell.border = bold_border

    # Reverse Polarity criteria
    cell = ws.cell(row=criteria_row, column=rev_pol_col)
    cell.value = "Pass"
    cell.font = criteria_font
    cell.alignment = center
    cell.border = bold_border

    # Unloaded Current criteria (grey)
    cell = ws.cell(row=criteria_row, column=unloaded_col)
    cell.fill = grey_fill
    cell.border = bold_border

    # Outer bold border around entire report
    apply_outer_border(ws, f'B2:{get_column_letter(last_col)}{criteria_row}')

    # Save
    wb.save(output_path)